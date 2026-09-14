import openpyxl
import xlrd
import json
import os
import sys
import glob

# UTF-8 输出
sys.stdout.reconfigure(encoding='utf-8')

# 附件目录：优先使用当前会话，回退到上一个会话
ATTACH_DIRS = [
    r'c:\Users\UserX\.trae-cn\attachments\6a87e8b0a13254c164ad5ca3',
    r'c:\Users\UserX\.trae-cn\attachments\6a843424826437c145e3a3ef',
]

adir = None
for d in ATTACH_DIRS:
    if os.path.isdir(d) and os.listdir(d):
        adir = d
        break

if adir is None:
    print("错误：找不到包含数据的附件目录", file=sys.stderr)
    sys.exit(1)

print(f"使用附件目录: {adir}\n")

# 按关键词查找文件
def find_file(*keywords):
    for f in os.listdir(adir):
        if all(kw in f for kw in keywords):
            return os.path.join(adir, f)
    return None

# 输出结果收集
results = {}

# ========== 1. 减灾能力 ==========
path1 = find_file('减灾能力')
if not path1:
    print("警告：未找到减灾能力文件，跳过", file=sys.stderr)
else:
    try:
        wb1 = openpyxl.load_workbook(path1, data_only=True)
        ws1 = wb1[wb1.sheetnames[0]]
        headers = [c.value for c in ws1[1]]
        print("=== 减灾能力 表头 ===")
        for i, h in enumerate(headers):
            if h:
                print(f"  col {i}: {h}")

        rows_data = []
        print("\n=== 减灾能力 罗卜田 数据 ===")
        for row in ws1.iter_rows(min_row=3, max_row=ws1.max_row, values_only=True):
            name = str(row[2]) if row[2] else ''
            addr = str(row[3]) if row[3] else ''
            if '罗卜田' in name or '罗卜田' in addr:
                d = {}
                for i, h in enumerate(headers):
                    if h and row[i] is not None:
                        d[h] = row[i]
                rows_data.append(d)
                print(json.dumps(d, ensure_ascii=False, default=str))
        results['减灾能力'] = rows_data
        print(f"  共 {len(rows_data)} 条记录")
    except Exception as e:
        print(f"错误：读取减灾能力文件失败: {e}", file=sys.stderr)

# ========== 2. 重点部位 ==========
path2 = find_file('重点部位')
if not path2:
    print("\n警告：未找到重点部位文件，跳过", file=sys.stderr)
else:
    try:
        wb2 = openpyxl.load_workbook(path2, data_only=True)
        all_sheets = {}
        for sn in wb2.sheetnames:
            ws = wb2[sn]
            print(f"\n=== 重点部位 Sheet: {sn} (rows={ws.max_row}, cols={ws.max_column}) ===")
            sheet_data = []
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
                vals = [v for v in row if v is not None]
                if vals:
                    sheet_data.append(vals)
                    print(vals)
            all_sheets[sn] = sheet_data
        results['重点部位'] = all_sheets
    except Exception as e:
        print(f"错误：读取重点部位文件失败: {e}", file=sys.stderr)

# ========== 3. 供水工程 ==========
path3 = find_file('供水工程')
if not path3:
    print("\n警告：未找到供水工程文件，跳过", file=sys.stderr)
else:
    try:
        wb3 = xlrd.open_workbook(path3)
        ws3 = wb3.sheet_by_name(wb3.sheet_names()[0])
        print(f"\n=== 供水工程 表头 (row 3) ===")
        headers3 = []
        for c in range(ws3.ncols):
            v = ws3.cell_value(2, c)
            headers3.append(v)
            if v:
                print(f"  col {c}: {v}")

        rows_data = []
        print(f"\n=== 供水工程 罗卜田 数据 ===")
        for r in range(3, ws3.nrows):
            name = str(ws3.cell_value(r, 1))
            loc = str(ws3.cell_value(r, 3))
            if '罗卜田' in name or '罗卜田' in loc:
                d = {}
                for c in range(ws3.ncols):
                    h = headers3[c]
                    v = ws3.cell_value(r, c)
                    if h and v != '':
                        d[h] = v
                rows_data.append(d)
                print(json.dumps(d, ensure_ascii=False, default=str))
        results['供水工程'] = rows_data
        print(f"  共 {len(rows_data)} 条记录")
    except Exception as e:
        print(f"错误：读取供水工程文件失败: {e}", file=sys.stderr)

# ========== 4. 重点敲门 ==========
path4 = find_file('重点敲门')
if not path4:
    print("\n警告：未找到重点敲门文件，跳过", file=sys.stderr)
else:
    try:
        wb4 = xlrd.open_workbook(path4)
        ws4 = wb4.sheet_by_name(wb4.sheet_names()[0])
        print(f"\n=== 重点敲门 总行数: {ws4.nrows} ===")
        all_rows = []
        for r in range(ws4.nrows):
            vals = [(c, ws4.cell_value(r, c)) for c in range(ws4.ncols) if ws4.cell_value(r, c) != '']
            if vals:
                all_rows.append(vals)
                print(vals)
        results['重点敲门'] = {'total_rows': ws4.nrows, 'data': all_rows}
        print(f"  共 {len(all_rows)} 条有效记录")
    except Exception as e:
        print(f"错误：读取重点敲门文件失败: {e}", file=sys.stderr)

# ========== 5. 燃气 ==========
path5 = find_file('燃气', '百日')
if not path5:
    path5 = find_file('燃气')
if not path5:
    path5 = find_file('百日')
if not path5:
    print("\n警告：未找到燃气文件，跳过", file=sys.stderr)
else:
    try:
        wb5 = openpyxl.load_workbook(path5, data_only=True)
        ws5 = wb5[wb5.sheetnames[0]]
        print(f"\n=== 燃气 全部数据 (rows={ws5.max_row}) ===")
        all_rows = []
        for row in ws5.iter_rows(min_row=1, max_row=ws5.max_row, values_only=True):
            vals = [v for v in row if v is not None]
            if vals:
                all_rows.append(vals)
                print(vals)
        results['燃气'] = all_rows
        print(f"  共 {len(all_rows)} 条记录")
    except Exception as e:
        print(f"错误：读取燃气文件失败: {e}", file=sys.stderr)

# ========== 输出 JSON ==========
output_path = os.path.join(adir, 'extracted_data.json')
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(results, f, ensure_ascii=False, default=str, indent=2)
print(f"\n=== 结果已保存到: {output_path} ===")
print(f"=== 提取完成，共处理 {len(results)} 个数据源 ===")
